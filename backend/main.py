"""CareerMatch AI MVP API. Run: uvicorn main:app --reload"""
import io
import os
import re
import base64
import hashlib
import hmac
import json
import secrets
import sqlite3
import time
from functools import lru_cache
from typing import Annotated

import fitz
from docx import Document
from fastapi import FastAPI, File, Form, HTTPException, Request, Response, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, Field
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

app = FastAPI(title="CareerMatch AI", version="1.0.0")
app.add_middleware(CORSMiddleware, allow_origins=os.getenv("FRONTEND_ORIGINS", "http://localhost:5173,http://127.0.0.1:5173").split(","), allow_credentials=True, allow_methods=["GET", "POST"], allow_headers=["*"])
MAX_UPLOAD = 5 * 1024 * 1024
AUTH_DB_PATH = os.path.abspath(os.getenv("AUTH_DB_PATH", os.path.join(os.path.dirname(__file__), "careermatch.sqlite3")))
AUTH_SECRET = os.getenv("AUTH_SECRET")
AUTH_SECRET_BYTES = AUTH_SECRET.encode("utf-8") if AUTH_SECRET else secrets.token_bytes(32)
AUTH_COOKIE = "cm_session"
AUTH_SESSION_SECONDS = 60 * 60 * 24 * 7
AUTH_COOKIE_SECURE = os.getenv("AUTH_COOKIE_SECURE", "false").lower() == "true"
SKILLS = ["Python", "SQL", "Excel", "Power BI", "Tableau", "Java", "JavaScript", "TypeScript", "React", "Node.js", "FastAPI", "Django", "Spring Boot", "PostgreSQL", "MySQL", "MongoDB", "Git", "Docker", "AWS", "Azure", "Machine Learning", "Data Analysis", "Data Visualization", "Statistics", "Communication", "Problem Solving", "Project Management", "Agile", "Scrum", "Customer Service", "Marketing", "SEO", "Figma", "HTML", "CSS", "REST APIs", "Testing", "Accounting", "Budgeting", "Sales", "Leadership", "Research", "Microsoft Word", "Presentation", "Data Cleaning", "ETL", "Pandas", "NumPy", "Scikit-learn"]
ALIASES = {"power bi": ["powerbi"], "javascript": ["js"], "typescript": ["ts"], "postgresql": ["postgres"], "data visualization": ["data visualisation"], "machine learning": ["ml"], "rest apis": ["rest api"], "node.js": ["nodejs", "node js"], "scikit-learn": ["sklearn"]}

def detect_skills(text: str) -> list[str]:
    found = []
    for skill in SKILLS:
        terms = [skill.lower()] + ALIASES.get(skill.lower(), [])
        if any(re.search(r"(?<!\w)" + re.escape(term) + r"(?!\w)", text.lower()) for term in terms):
            found.append(skill)
    return found

@lru_cache(maxsize=1)
def embedding_model():
    from sentence_transformers import SentenceTransformer
    return SentenceTransformer("sentence-transformers/all-MiniLM-L6-v2")

def similarity(a: str, b: str) -> tuple[float, str]:
    if os.getenv("MATCH_ENGINE", "semantic").lower() == "tfidf":
        engine = "tfidf"
    else:
        engine = "semantic"
    if engine == "semantic":
        try:
            model = embedding_model()
            vectors = model.encode([a[:12000], b[:12000]], normalize_embeddings=True)
            return max(0.0, min(1.0, float(vectors[0] @ vectors[1]))), "sentence-transformers"
        except Exception as exc:
            print(f"Semantic model unavailable, using TF-IDF: {exc}")
    vectors = TfidfVectorizer(stop_words="english").fit_transform([a[:12000], b[:12000]])
    return max(0.0, min(1.0, float(cosine_similarity(vectors[0], vectors[1])[0][0]))), "tfidf-fallback"

class MatchRequest(BaseModel):
    cv_text: str = Field(min_length=30, max_length=100000)
    job_description: str = Field(min_length=30, max_length=30000)
    job_title: str = Field(default="Target role", max_length=120)

class RegisterRequest(BaseModel):
    name: str = Field(min_length=1, max_length=100)
    email: str = Field(min_length=3, max_length=254)
    password: str = Field(min_length=8, max_length=128)

class LoginRequest(BaseModel):
    email: str = Field(min_length=3, max_length=254)
    password: str = Field(min_length=1, max_length=128)

def initialize_auth_db():
    os.makedirs(os.path.dirname(AUTH_DB_PATH), exist_ok=True)
    with sqlite3.connect(AUTH_DB_PATH) as connection:
        connection.execute("""CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            email TEXT NOT NULL UNIQUE COLLATE NOCASE,
            password_salt TEXT NOT NULL,
            password_hash TEXT NOT NULL,
            created_at INTEGER NOT NULL
        )""")

def _password_hash(password: str, salt: bytes) -> bytes:
    return hashlib.scrypt(password.encode("utf-8"), salt=salt, n=2**14, r=8, p=1, dklen=64)

def _encode_token_part(value: bytes) -> str:
    return base64.urlsafe_b64encode(value).rstrip(b"=").decode("ascii")

def _create_session_token(user_id: int) -> str:
    payload = _encode_token_part(json.dumps({"uid": user_id, "exp": int(time.time()) + AUTH_SESSION_SECONDS}).encode("utf-8"))
    signature = _encode_token_part(hmac.new(AUTH_SECRET_BYTES, payload.encode("ascii"), hashlib.sha256).digest())
    return f"{payload}.{signature}"

def _read_session_user(request: Request) -> dict:
    token = request.cookies.get(AUTH_COOKIE, "")
    try:
        payload_part, signature = token.split(".", 1)
        expected = _encode_token_part(hmac.new(AUTH_SECRET_BYTES, payload_part.encode("ascii"), hashlib.sha256).digest())
        if not hmac.compare_digest(signature, expected):
            raise ValueError("Invalid signature")
        padded_payload = payload_part + "=" * (-len(payload_part) % 4)
        payload = json.loads(base64.urlsafe_b64decode(padded_payload.encode("ascii")))
        if int(payload["exp"]) < int(time.time()):
            raise ValueError("Expired session")
        user_id = int(payload["uid"])
    except (ValueError, KeyError, TypeError, json.JSONDecodeError):
        raise HTTPException(401, "Please sign in to continue.")
    with sqlite3.connect(AUTH_DB_PATH) as connection:
        connection.row_factory = sqlite3.Row
        user = connection.execute("SELECT id, name, email FROM users WHERE id = ?", (user_id,)).fetchone()
    if user is None:
        raise HTTPException(401, "Please sign in to continue.")
    return dict(user)

def _set_session_cookie(response: Response, user_id: int):
    response.set_cookie(AUTH_COOKIE, _create_session_token(user_id), max_age=AUTH_SESSION_SECONDS, httponly=True, secure=AUTH_COOKIE_SECURE, samesite="lax", path="/")

initialize_auth_db()

@app.post("/api/auth/register", status_code=201)
def register(payload: RegisterRequest, response: Response):
    name = payload.name.strip()
    email = payload.email.strip().lower()
    if not name:
        raise HTTPException(422, "Enter your name.")
    if not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):
        raise HTTPException(422, "Enter a valid email address.")
    salt = secrets.token_bytes(16)
    password_hash = _password_hash(payload.password, salt)
    try:
        with sqlite3.connect(AUTH_DB_PATH) as connection:
            cursor = connection.execute("INSERT INTO users (name, email, password_salt, password_hash, created_at) VALUES (?, ?, ?, ?, ?)", (name, email, salt.hex(), password_hash.hex(), int(time.time())))
            user_id = cursor.lastrowid
    except sqlite3.IntegrityError:
        raise HTTPException(409, "An account with this email already exists.")
    _set_session_cookie(response, user_id)
    return {"user": {"id": user_id, "name": name, "email": email}}

@app.post("/api/auth/login")
def login(payload: LoginRequest, response: Response):
    email = payload.email.strip().lower()
    with sqlite3.connect(AUTH_DB_PATH) as connection:
        connection.row_factory = sqlite3.Row
        user = connection.execute("SELECT id, name, email, password_salt, password_hash FROM users WHERE email = ?", (email,)).fetchone()
    if user is None:
        raise HTTPException(401, "Email or password is incorrect.")
    salt = bytes.fromhex(user["password_salt"])
    if not hmac.compare_digest(_password_hash(payload.password, salt).hex(), user["password_hash"]):
        raise HTTPException(401, "Email or password is incorrect.")
    _set_session_cookie(response, user["id"])
    return {"user": {"id": user["id"], "name": user["name"], "email": user["email"]}}

@app.get("/api/auth/me")
def current_user(request: Request):
    return {"user": _read_session_user(request)}

@app.post("/api/auth/logout")
def logout(response: Response):
    response.delete_cookie(AUTH_COOKIE, httponly=True, secure=AUTH_COOKIE_SECURE, samesite="lax", path="/")
    return {"status": "signed_out"}

@app.get("/api/health")
def health():
    return {"status": "ok", "service": "CareerMatch AI"}

@app.post("/api/parse-cv")
async def parse_cv(file: Annotated[UploadFile, File()]):
    filename = file.filename or ""
    ext = os.path.splitext(filename)[1].lower()
    if ext not in (".pdf", ".docx", ".txt"):
        raise HTTPException(400, "Upload a PDF, DOCX or TXT file.")
    data = await file.read(MAX_UPLOAD + 1)
    if len(data) > MAX_UPLOAD:
        raise HTTPException(413, "File exceeds the 5 MB limit.")
    try:
        if ext == ".pdf":
            with fitz.open(stream=data, filetype="pdf") as pdf:
                text = "\n".join(page.get_text() for page in pdf)
        elif ext == ".docx":
            document = Document(io.BytesIO(data))
            text = "\n".join([p.text for p in document.paragraphs] + [cell.text for table in document.tables for row in table.rows for cell in row.cells])
        else:
            text = data.decode("utf-8-sig")
    except Exception:
        raise HTTPException(422, "Could not read the document. Try a text-based PDF, DOCX or TXT file.")
    if len(text.strip()) < 30:
        raise HTTPException(422, "Not enough selectable text found. Scanned/image-only PDFs are not supported yet.")
    return {"filename": filename, "text": text[:100000], "skills": detect_skills(text)}

@app.post("/api/match")
def match(payload: MatchRequest):
    cv_skills = detect_skills(payload.cv_text)
    job_skills = detect_skills(payload.job_description)
    matched = [s for s in job_skills if s in cv_skills]
    missing = [s for s in job_skills if s not in cv_skills]
    semantic, engine = similarity(payload.cv_text, payload.job_description)
    # Skill coverage is 65%, text similarity is 35%; if no skills found, only similarity is used.
    coverage = len(matched) / len(job_skills) if job_skills else None
    score = round(100 * (0.65 * coverage + 0.35 * semantic)) if coverage is not None else round(100 * semantic)
    score = max(0, min(100, score))
    if missing:
        recommendations = [f"Build or demonstrate practical experience with {s}." for s in missing[:5]]
    else:
        recommendations = ["Tailor your CV bullet points to the responsibilities in this job description.", "Use measurable examples to demonstrate relevant achievements."]
    recommendations.append("Check that your CV clearly describes your actual experience; never add skills you do not have.")
    return {"job_title": payload.job_title, "score": score, "matched_skills": matched, "missing_skills": missing, "cv_skills": cv_skills, "job_skills": job_skills, "skill_coverage": round(100 * coverage) if coverage is not None else None, "semantic_similarity": round(100 * semantic), "engine": engine, "recommendations": recommendations, "disclaimer": "This is an indicative text-to-job fit score, not a hiring probability or hiring decision."}
